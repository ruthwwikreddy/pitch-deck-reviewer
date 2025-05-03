import streamlit as st
import fitz  # PyMuPDF
from pptx import Presentation
import subprocess
import json
import tempfile
import os

# -------------------- Text Extraction --------------------

def extract_text_from_pdf(file):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
        tmp_file.write(file.read())
        tmp_path = tmp_file.name
    doc = fitz.open(tmp_path)
    text = "\n".join(page.get_text() for page in doc)
    os.remove(tmp_path)
    return text

def extract_text_from_pptx(file):
    prs = Presentation(file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text.strip() + "\n"
    return text

# -------------------- Ollama Interaction --------------------

def call_model(text, model="mistral:latest"):
    prompt = f"""
You are a professional startup pitch deck reviewer.
Evaluate this content and return a JSON like this:
{{
  "scores": {{
    "Clarity": {{ "score": 4, "reason": "..." }},
    "Problem": {{ "score": 5, "reason": "..." }},
    "Solution": {{ "score": 4, "reason": "..." }},
    "Market": {{ "score": 3, "reason": "..." }},
    "Team": {{ "score": 4, "reason": "..." }},
    "Design": {{ "score": 4, "reason": "..." }}
  }},
  "summary": "Overall feedback summary..."
}}

Pitch Deck Content:
{text[:6000]}
"""
    result = subprocess.run(
        ["ollama", "run", model],
        input=prompt.encode(),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        timeout=120
    )
    return result.stdout.decode()

# -------------------- Star Rendering --------------------

def render_star_rating(score):
    return "⭐" * int(score) + "☆" * (5 - int(score))

# -------------------- Streamlit UI --------------------

st.set_page_config("Pitch Deck Reviewer", layout="centered")
st.title("Pitch Deck Reviewer")
st.markdown("Upload your deck, choose a model, and let AI review it with detailed feedback.")

uploaded_file = st.file_uploader("Upload your pitch deck (.pdf / .pptx)", type=["pdf", "pptx"])

model_choice = st.selectbox("Choose AI Model", [
    "mistral:latest", "phi3:latest", "llama3:latest", "deepseek-llm:latest"
])

if uploaded_file and model_choice:
    if uploaded_file.name.endswith(".pdf"):
        file_text = extract_text_from_pdf(uploaded_file)
    else:
        file_text = extract_text_from_pptx(uploaded_file)

    if st.button("Analyze Deck"):
        with st.spinner(f"Reviewing with {model_choice}..."):
            try:
                response = call_model(file_text, model=model_choice)
                parsed = json.loads(response)

                st.markdown("## Review Results")

                for cat, data in parsed["scores"].items():
                    score = min(max(int(data["score"]), 0), 5)  # Normalize
                    st.markdown(f"**{cat}**: {render_star_rating(score)} ({score}/5)")
                    st.caption(f"{data['reason']}")

                st.markdown("## Summary")
                st.success(parsed["summary"])

                report = "# AI Pitch Deck Review Report\n\n"
                for k, v in parsed["scores"].items():
                    report += f"## {k}\n- **Score**: {v['score']} / 5\n- **Reason**: {v['reason']}\n\n"
                report += f"## Summary\n{parsed['summary']}\n"

                st.download_button("📥 Download Full Report", report, file_name="pitch_deck_review.md")

            except Exception as e:
                st.error("⚠️ Review failed or response couldn't be parsed.")
                st.text(response)
